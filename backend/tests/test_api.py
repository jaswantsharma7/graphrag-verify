"""API user flows against temporary durable storage and the real demo pipeline."""

import asyncio
import io
import tempfile
import threading
import time
import unittest
from pathlib import Path
from unittest.mock import patch

from fastapi.testclient import TestClient

from app.config import Settings
from app.ingestion import SourceStorage
from app.main import create_app
from app.pipeline import Pipeline
from app.repository import Repository


TEXT_A = "Aster Holdings owns Beacon Components. Beacon Components supplies Cedar Robotics."
TEXT_B = "Orion Labs owns Private Systems. Private Systems supplies Hidden Robotics."


def pdf_bytes(texts: list[str]) -> bytes:
    from pypdf import PdfWriter
    from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

    writer = PdfWriter()
    for text in texts:
        page = writer.add_blank_page(width=612, height=792)
        font = DictionaryObject({
            NameObject("/Type"): NameObject("/Font"),
            NameObject("/Subtype"): NameObject("/Type1"),
            NameObject("/BaseFont"): NameObject("/Helvetica"),
        })
        page[NameObject("/Resources")] = DictionaryObject({
            NameObject("/Font"): DictionaryObject({NameObject("/F1"): writer._add_object(font)}),
        })
        stream = DecodedStreamObject()
        stream.set_data(f"BT /F1 12 Tf 50 740 Td ({text}) Tj ET".encode("ascii"))
        page[NameObject("/Contents")] = writer._add_object(stream)
    output = io.BytesIO()
    writer.write(output)
    return output.getvalue()


class APITests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)
        self.settings = Settings(
            _env_file=None, app_mode="demo", data_dir=self.root / "data",
            api_key=None, embedding_dimensions=128,
        )

    def tearDown(self):
        self.temp.cleanup()

    def wait_job(self, client, job_id, expected="completed"):
        deadline = time.monotonic() + 10
        while time.monotonic() < deadline:
            response = client.get(f"/api/jobs/{job_id}")
            self.assertEqual(response.status_code, 200, response.text)
            job = response.json()
            self.assertGreaterEqual(job["progress"], 0)
            self.assertLessEqual(job["progress"], 1)
            if job["status"] in {"completed", "failed"}:
                self.assertEqual(job["status"], expected, job)
                return job
            time.sleep(0.02)
        self.fail("Ingestion did not finish within the test deadline.")

    def add_text(self, client, name, text):
        response = client.post("/api/documents/text", json={"name": name, "text": text})
        self.assertEqual(response.status_code, 202, response.text)
        result = response.json()
        self.wait_job(client, result["job"]["id"])
        return result["document"]["id"]

    def test_ingestion_readback_dedupe_jobs_history_and_restart(self):
        with TestClient(create_app(self.settings)) as client:
            self.assertEqual(client.get("/api/documents").json(), {"documents": []})
            self.assertEqual(client.get("/api/health").json()["status"], "ok")
            document_id = self.add_text(client, "Supply chain", TEXT_A)
            document = client.get("/api/documents").json()["documents"][0]
            self.assertEqual(document["status"], "ready")
            self.assertGreater(document["chunks"], 0)
            self.assertIsNone(document["error"])
            self.assertNotIn("storage_name", document)
            duplicate = client.post(
                "/api/documents", files={"file": ("renamed.md", TEXT_A.encode(), "text/markdown")},
            )
            self.assertEqual(duplicate.status_code, 202)
            self.assertEqual(duplicate.json()["document"]["id"], document_id)
            self.assertEqual(duplicate.json()["job"]["id"], document["job_id"])
            self.assertEqual(len(list((self.settings.data_dir / "uploads").iterdir())), 1)
            result = client.post("/api/query", json={
                "question": "Who does Beacon Components supply?", "document_ids": [document_id],
            })
            self.assertEqual(result.status_code, 200, result.text)
            query = result.json()
            self.assertEqual(query["mode"], "demo")
            self.assertTrue(query["evidence"])
            self.assertEqual({item["document_id"] for item in query["evidence"]}, {document_id})
            self.assertIn("Cedar Robotics", query["answer"])
            self.assertEqual(client.get("/api/queries").json()["queries"][0]["id"], query["id"])
            self.assertEqual(client.get("/api/jobs/not-found").status_code, 404)
        with TestClient(create_app(self.settings)) as client:
            self.assertEqual(client.get("/api/documents").json()["documents"][0]["id"], document_id)
            self.assertEqual(client.get("/api/queries").json()["queries"][0]["id"], query["id"])
            result = client.post("/api/query", json={"question": "Who supplies Cedar Robotics?"})
            self.assertEqual(result.status_code, 200)
            self.assertTrue(result.json()["evidence"])

    def test_filters_scope_query_graph_and_deletion_removes_sources(self):
        with TestClient(create_app(self.settings)) as client:
            public_id = self.add_text(client, "Public", TEXT_A)
            private_id = self.add_text(client, "Private", TEXT_B)
            result = client.post("/api/query", json={
                "question": "Who does Aster Holdings own?", "document_ids": [public_id],
            }).json()
            self.assertTrue(result["evidence"])
            self.assertEqual({item["document_id"] for item in result["evidence"]}, {public_id})
            self.assertNotIn("Private Systems", result["answer"])
            graph = client.get("/api/graph", params={"document_ids": public_id, "limit": 500})
            self.assertEqual(graph.status_code, 200, graph.text)
            self.assertTrue(graph.json()["nodes"])
            self.assertNotIn("Private Systems", {node["label"] for node in graph.json()["nodes"]})
            for edge in graph.json()["edges"]:
                self.assertTrue(all(identifier.startswith(public_id) for identifier in edge["evidence_ids"]))
            self.assertEqual(client.post("/api/query", json={
                "question": "Ownership?", "document_ids": ["missing"],
            }).status_code, 404)
            self.assertEqual(client.get("/api/graph?document_ids=missing").status_code, 404)
            repository = Repository(self.settings.data_dir)
            storage_name = repository.get_document_record(public_id)["storage_name"]
            job_id = repository.get_document(public_id)["job_id"]
            self.assertEqual(client.delete(f"/api/documents/{public_id}").status_code, 204)
            self.assertFalse((self.settings.data_dir / "uploads" / storage_name).exists())
            self.assertEqual(client.get(f"/api/jobs/{job_id}").status_code, 404)
            self.assertEqual(client.delete(f"/api/documents/{public_id}").status_code, 404)
            self.assertEqual(client.get("/api/graph", params={"document_ids": public_id}).status_code, 404)
            self.assertEqual(client.delete(f"/api/documents/{private_id}").status_code, 204)
            self.assertEqual(client.get("/api/graph").json(), {"nodes": [], "edges": []})
            self.assertEqual(client.get("/api/health").json()["stats"], {
                "documents": 0, "chunks": 0, "entities": 0, "relationships": 0,
            })
            empty = client.post("/api/query", json={"question": "Who owns Beacon Components?"}).json()
            self.assertEqual(empty["status"], "no_evidence")
            self.assertEqual(empty["evidence"], [])

    def test_failed_job_is_safe_and_only_explicit_retry_requeues(self):
        pipeline = Pipeline(self.settings)

        async def unavailable(*args):
            raise RuntimeError("private-provider-token-and-source-text")

        with TestClient(create_app(self.settings, pipeline=pipeline)) as client:
            with patch.object(pipeline, "ingest", new=unavailable):
                uploaded = client.post("/api/documents/text", json={"name": "Retry me", "text": TEXT_A}).json()
                document_id, old_job = uploaded["document"]["id"], uploaded["job"]["id"]
                job = self.wait_job(client, old_job, "failed")
            self.assertIn("retry", job["error"].lower())
            self.assertNotIn("private-provider", job["error"])
            failed = client.get("/api/documents").json()["documents"][0]
            self.assertEqual(failed["status"], "failed")
            self.assertEqual(failed["chunks"], 0)
            self.assertEqual(failed["error"], job["error"])
            duplicate = client.post("/api/documents/text", json={"name": "Duplicate", "text": TEXT_A}).json()
            self.assertEqual(duplicate["job"]["id"], old_job)
            self.assertEqual(duplicate["job"]["status"], "failed")
            self.assertEqual(client.post("/api/query", json={
                "question": "Who owns Beacon?", "document_ids": [document_id],
            }).status_code, 409)
            self.assertEqual(client.get("/api/graph", params={"document_ids": document_id}).status_code, 409)
            self.assertEqual(client.post("/api/query", json={"question": "Who owns Beacon?"}).json()["evidence"], [])
            retry = client.post(f"/api/documents/{document_id}/retry")
            self.assertEqual(retry.status_code, 202, retry.text)
            self.assertNotEqual(retry.json()["job"]["id"], old_job)
            self.wait_job(client, retry.json()["job"]["id"])
            self.assertEqual(client.get(f"/api/jobs/{old_job}").json()["status"], "failed")
            self.assertEqual(client.post(f"/api/documents/{document_id}/retry").status_code, 409)
            self.assertEqual(client.post("/api/documents/missing/retry").status_code, 404)
            result = client.post("/api/query", json={"question": "Who owns Beacon Components?"}).json()
            self.assertTrue(result["evidence"])

    def test_active_worker_keeps_second_job_queued_and_blocks_filters_and_delete(self):
        pipeline, entered, release = Pipeline(self.settings), threading.Event(), threading.Event()
        original = pipeline.ingest

        async def blocked(document_id, name, chunks):
            entered.set()
            while not release.is_set():
                await asyncio.sleep(0.01)
            return await original(document_id, name, chunks)

        with TestClient(create_app(self.settings, pipeline=pipeline)) as client:
            with patch.object(pipeline, "ingest", new=blocked):
                first = client.post("/api/documents/text", json={"name": "First", "text": TEXT_A}).json()
                try:
                    self.assertTrue(entered.wait(timeout=3))
                    second = client.post("/api/documents/text", json={"name": "Second", "text": TEXT_B}).json()
                    self.assertEqual(client.get(f"/api/jobs/{first['job']['id']}").json()["status"], "processing")
                    self.assertEqual(client.get(f"/api/jobs/{second['job']['id']}").json()["status"], "queued")
                    for result in (first, second):
                        document_id = result["document"]["id"]
                        self.assertEqual(client.delete(f"/api/documents/{document_id}").status_code, 409)
                        self.assertEqual(client.post("/api/query", json={
                            "question": "Who owns Beacon?", "document_ids": [document_id],
                        }).status_code, 409)
                        self.assertEqual(client.get("/api/graph", params={"document_ids": document_id}).status_code, 409)
                finally:
                    release.set()
                self.wait_job(client, first["job"]["id"])
                self.wait_job(client, second["job"]["id"])

    def test_interrupted_processing_recovers_to_failed_then_retries(self):
        repository = Repository(self.settings.data_dir)
        storage = SourceStorage(self.settings.data_dir, self.settings.max_upload_mb)
        storage_name, digest = storage.text(TEXT_A)
        document, job, _ = repository.create_document("Interrupted", digest, storage_name)
        self.assertEqual(repository.claim_next_job()[1]["id"], job["id"])
        self.assertIsNone(repository.claim_next_job())
        with TestClient(create_app(self.settings)) as client:
            failed_job = client.get(f"/api/jobs/{job['id']}").json()
            self.assertEqual(failed_job["status"], "failed")
            self.assertIn("interrupted", failed_job["error"].lower())
            failed_document = client.get("/api/documents").json()["documents"][0]
            self.assertEqual(failed_document["status"], "failed")
            self.assertEqual(failed_document["error"], failed_job["error"])
            retry = client.post(f"/api/documents/{document['id']}/retry")
            self.assertEqual(retry.status_code, 202)
            self.wait_job(client, retry.json()["job"]["id"])
            self.assertEqual(client.get("/api/documents").json()["documents"][0]["status"], "ready")

    def test_optional_auth_protects_all_api_and_liveness_remains_public(self):
        settings = self.settings.model_copy(update={"api_key": "unit-test-key"})
        with TestClient(create_app(settings)) as client:
            self.assertEqual(client.get("/healthz").status_code, 200)
            for path in ("/api/health", "/api/documents", "/api/queries", "/api/graph", "/api/unknown"):
                self.assertEqual(client.get(path).status_code, 401)
                self.assertEqual(client.get(path, headers={"X-API-Key": "wrong"}).status_code, 401)
            self.assertEqual(client.post("/api/documents/text", json={"name": "Blocked", "text": TEXT_A}).status_code, 401)
            self.assertEqual(client.delete("/api/documents/missing").status_code, 401)
            headers = {"X-API-Key": "unit-test-key"}
            self.assertEqual(client.get("/api/health", headers=headers).status_code, 200)
            self.assertEqual(client.get("/api/unknown", headers=headers).status_code, 404)

    def test_absent_live_configuration_reports_degraded_without_startup_crash(self):
        settings = self.settings.model_copy(update={
            "app_mode": "live", "neo4j_password": "", "llm_api_key": "", "embedding_api_key": "",
        })
        with TestClient(create_app(settings)) as client:
            self.assertEqual(client.get("/healthz").status_code, 200)
            health = client.get("/api/health")
            self.assertEqual(health.status_code, 200)
            self.assertEqual(health.json()["status"], "degraded")
            self.assertFalse(health.json()["pipeline_ready"])
            self.assertNotIn("neo4j_password", health.text)
            self.assertEqual(client.get("/api/documents").json(), {"documents": []})
            self.assertEqual(client.post("/api/query", json={"question": "Ownership?"}).status_code, 503)
            self.assertEqual(client.post("/api/demo/seed").status_code, 409)

    def test_upload_and_text_validation_and_bounded_request_sizes(self):
        settings = self.settings.model_copy(update={"max_upload_mb": 1, "max_document_chars": 100})
        with TestClient(create_app(settings)) as client:
            for payload, status in (
                ({"name": "../escape", "text": "valid"}, 400),
                ({"name": "a\nb", "text": "valid"}, 400),
                ({"name": "Empty", "text": "   "}, 400),
                ({"name": "Invalid text", "text": "has\x00null"}, 400),
                ({"name": "Large", "text": "x" * 101}, 413),
                ({"name": "Missing"}, 422),
                ({"name": "Extra", "text": "valid", "unrecognized": True}, 422),
            ):
                response = client.post("/api/documents/text", json=payload)
                self.assertEqual(response.status_code, status, response.text)
                self.assertIsInstance(response.json()["detail"], str)
            self.assertEqual(client.post("/api/documents", files={"file": ("bad.exe", b"text")}).status_code, 400)
            self.assertEqual(client.post("/api/documents", files={"file": ("empty.txt", b"")}).status_code, 400)
            self.assertEqual(client.post("/api/documents", files={"file": ("../bad.txt", b"text")}).status_code, 400)
            oversized = client.post("/api/documents", files={"file": ("large.txt", b"x" * (1024 * 1024 + 1))})
            self.assertEqual(oversized.status_code, 413)
            streamed = client.post(
                "/api/documents/text", content=iter([b'{"name":"Bounded","text":"', b"x" * 16000, b'"}']),
                headers={"Content-Type": "application/json"},
            )
            self.assertEqual(streamed.status_code, 413, streamed.text)
            invalid = client.post("/api/query", json={"question": "private prompt content", "top_k": 100})
            self.assertEqual(invalid.status_code, 422)
            self.assertNotIn("private prompt content", invalid.text)
            self.assertNotIn("Traceback", invalid.text)
            self.assertEqual(client.get("/api/documents").json()["documents"], [])
            self.assertEqual(list((settings.data_dir / "uploads").iterdir()), [])

    def test_pdf_physical_pages_docx_and_utf8_markdown(self):
        from docx import Document

        with TestClient(create_app(self.settings)) as client:
            source = pdf_bytes([
                "",
                "Aster Holdings owns Beacon Components.",
                "Beacon Components supplies Cedar Robotics.",
            ])
            uploaded = client.post("/api/documents", files={"file": ("Report.PDF", source, "application/pdf")})
            self.assertEqual(uploaded.status_code, 202, uploaded.text)
            record = uploaded.json()
            self.wait_job(client, record["job"]["id"])
            query = client.post("/api/query", json={
                "question": "How does Aster Holdings connect to Cedar Robotics?",
                "document_ids": [record["document"]["id"]],
            })
            self.assertEqual(query.status_code, 200, query.text)
            self.assertEqual({item["page"] for item in query.json()["evidence"]}, {2, 3})
            self.assertEqual({item["document_name"] for item in query.json()["evidence"]}, {"Report.PDF"})
            word_document = Document()
            word_document.add_paragraph("Orion Labs owns Private Systems.")
            word_document.add_table(rows=1, cols=1).cell(0, 0).text = "Private Systems supplies Hidden Robotics."
            output = io.BytesIO()
            word_document.save(output)
            uploaded = client.post("/api/documents", files={"file": ("Notes.docx", output.getvalue())}).json()
            self.wait_job(client, uploaded["job"]["id"])
            query = client.post("/api/query", json={
                "question": "Who supplies Hidden Robotics?", "document_ids": [uploaded["document"]["id"]],
            }).json()
            self.assertTrue(query["evidence"])
            self.assertTrue(all(item["page"] is None for item in query["evidence"]))
            self.assertIn("Hidden Robotics", query["answer"])
            markdown = client.post("/api/documents", files={"file": (
                "Unicode.md", "# Notes\nCafé Research uses Delta Systems.".encode("utf-8"), "text/markdown",
            )}).json()
            self.wait_job(client, markdown["job"]["id"])
            query = client.post("/api/query", json={
                "question": "What uses Delta Systems?", "document_ids": [markdown["document"]["id"]],
            }).json()
            self.assertTrue(query["evidence"])
            self.assertIn("Café", query["evidence"][0]["text"])
            malformed = client.post("/api/documents", files={"file": ("Invalid.txt", b"\xff\xfe")}).json()
            failed = self.wait_job(client, malformed["job"]["id"], "failed")
            self.assertIn("UTF-8", failed["error"])
            self.assertEqual(client.delete(f"/api/documents/{malformed['document']['id']}").status_code, 204)

    def test_no_text_pdf_and_chunk_limits_fail_without_indexed_sources(self):
        from pypdf import PdfWriter

        writer, output = PdfWriter(), io.BytesIO()
        writer.add_blank_page(width=612, height=792)
        writer.write(output)
        settings = self.settings.model_copy(update={"chunk_size": 64, "chunk_overlap": 0, "max_chunks_per_document": 1})
        with TestClient(create_app(settings)) as client:
            uploaded = client.post("/api/documents", files={"file": ("NoText.pdf", output.getvalue())}).json()
            failed = self.wait_job(client, uploaded["job"]["id"], "failed")
            self.assertIn("OCR", failed["error"])
            large = client.post("/api/documents/text", json={"name": "Too many chunks", "text": TEXT_A * 80}).json()
            failed = self.wait_job(client, large["job"]["id"], "failed")
            self.assertIn("chunk limit", failed["error"])
            self.assertEqual(client.get("/api/health").json()["stats"]["chunks"], 0)
            self.assertEqual(client.post("/api/query", json={"question": "Who supplies Cedar?"}).json()["status"], "no_evidence")

    def test_query_history_is_newest_first_and_capped_at_one_hundred(self):
        with TestClient(create_app(self.settings)) as client:
            for number in range(105):
                response = client.post("/api/query", json={"question": f"Question {number}"})
                self.assertEqual(response.status_code, 200)
                self.assertEqual(response.json()["status"], "no_evidence")
            history = client.get("/api/queries?limit=500").json()["queries"]
            self.assertEqual(len(history), 100)
            self.assertEqual(history[0]["question"], "Question 104")
            self.assertEqual(history[-1]["question"], "Question 5")
            self.assertEqual(len(client.get("/api/queries").json()["queries"]), 20)
            self.assertEqual(client.get("/api/queries?limit=0").status_code, 422)

    def test_demo_seeding_is_explicit_and_idempotent(self):
        sample_root = self.root / "fixture-root"
        samples = sample_root / "samples"
        samples.mkdir(parents=True)
        (samples / "chain.txt").write_text(TEXT_A, encoding="utf-8")
        (samples / "private.md").write_text(TEXT_B, encoding="utf-8")
        (samples / "ignored.bin").write_bytes(b"not a supported sample")
        with patch("app.api.PROJECT_ROOT", sample_root):
            with TestClient(create_app(self.settings)) as client:
                self.assertEqual(client.get("/api/documents").json()["documents"], [])
                first = client.post("/api/demo/seed")
                self.assertEqual(first.status_code, 200, first.text)
                self.assertEqual(len(first.json()["documents"]), 2)
                for job in first.json()["jobs"]:
                    self.wait_job(client, job["id"])
                second = client.post("/api/demo/seed").json()
                self.assertEqual(
                    [item["id"] for item in first.json()["documents"]],
                    [item["id"] for item in second["documents"]],
                )
                self.assertEqual(
                    [item["id"] for item in first.json()["jobs"]],
                    [item["id"] for item in second["jobs"]],
                )
                self.assertEqual(len(list((self.settings.data_dir / "uploads").iterdir())), 2)

    def test_spa_fallback_serves_build_without_masking_api_or_exposing_sources(self):
        dist = self.root / "dist"
        dist.mkdir()
        (dist / "index.html").write_text("<!doctype html><title>Fixture frontend</title>", encoding="utf-8")
        (dist / "app.js").write_text("window.fixture = true;", encoding="utf-8")
        (self.root / "outside.txt").write_text("outside frontend", encoding="utf-8")
        settings = self.settings.model_copy(update={"data_dir": dist / "private"})
        with TestClient(create_app(settings, frontend_dist=dist)) as client:
            self.assertIn("Fixture frontend", client.get("/").text)
            self.assertIn("Fixture frontend", client.get("/documents/view").text)
            self.assertEqual(client.get("/app.js").status_code, 200)
            self.assertEqual(client.get("/missing.js").status_code, 404)
            api_missing = client.get("/api/not-a-route")
            self.assertEqual(api_missing.status_code, 404)
            self.assertIn("detail", api_missing.json())
            self.assertEqual(client.get("/%2e%2e/outside.txt").status_code, 404)
            document_id = self.add_text(client, "Private source", TEXT_A)
            source_name = Repository(settings.data_dir).get_document_record(document_id)["storage_name"]
            self.assertEqual(client.get(f"/private/uploads/{source_name}").status_code, 404)
            self.assertEqual(client.get("/private/metadata.sqlite3").status_code, 404)

    def test_shutdown_cancels_stalled_ingestion_and_keeps_retryable_failed_metadata(self):
        pipeline, entered = Pipeline(self.settings), threading.Event()

        async def stalled(*args):
            entered.set()
            await asyncio.Event().wait()

        with patch.object(pipeline, "ingest", new=stalled):
            with TestClient(create_app(self.settings, pipeline=pipeline)) as client:
                result = client.post("/api/documents/text", json={"name": "Interrupted shutdown", "text": TEXT_A}).json()
                self.assertTrue(entered.wait(timeout=3))
        with TestClient(create_app(self.settings)) as client:
            job = client.get(f"/api/jobs/{result['job']['id']}").json()
            self.assertEqual(job["status"], "failed")
            self.assertIn("shutdown", job["error"])
            document = client.get("/api/documents").json()["documents"][0]
            self.assertEqual(document["status"], "failed")
            self.assertEqual(document["error"], job["error"])
            retry = client.post(f"/api/documents/{document['id']}/retry")
            self.assertEqual(retry.status_code, 202)
            self.wait_job(client, retry.json()["job"]["id"])


if __name__ == "__main__":
    unittest.main()
